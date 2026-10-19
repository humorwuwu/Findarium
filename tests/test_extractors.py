from __future__ import annotations

from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from findarium.indexing.extractors import extract_document


def test_docx_extracts_heading_and_table(tmp_path: Path):
    path = tmp_path / "contract.docx"
    document = Document()
    document.add_heading("采购合同", level=1)
    document.add_paragraph("付款条款和验收安排。")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "金额"
    table.rows[0].cells[1].text = "1000"
    document.save(str(path))

    result = extract_document(path)
    assert result.title
    assert any(block.kind == "heading" for block in result.blocks)
    assert any(block.kind == "table" for block in result.blocks)
    assert "付款条款" in result.full_text


def test_xlsx_extracts_sheet_rows(tmp_path: Path):
    path = tmp_path / "budget.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "预算"
    sheet.append(["项目", "金额"])
    sheet.append(["采购", 1000])
    workbook.save(str(path))

    result = extract_document(path)
    assert result.sheet_count == 1
    assert any(block.location_json.get("sheet") == "预算" for block in result.blocks)
    assert "采购" in result.full_text


def test_pptx_extracts_slide(tmp_path: Path):
    path = tmp_path / "deck.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "项目汇报"
    slide.placeholders[1].text = "合同付款进度"
    presentation.save(str(path))

    result = extract_document(path)
    assert result.page_count == 1
    assert any(block.location_json.get("slide") == 1 for block in result.blocks)
    assert "合同付款" in result.full_text


def test_pdf_extracts_page_text(tmp_path: Path):
    path = tmp_path / "report.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "acceptance report 2025")
    document.save(str(path))
    document.close()

    result = extract_document(path)
    assert result.page_count == 1
    assert result.blocks[0].location_json["page"] == 1
    assert "acceptance report" in result.full_text


def test_markdown_extracts_links_and_headings(tmp_path: Path):
    path = tmp_path / "note.md"
    path.write_text("# 项目\n\n参见 [预算](budget.xlsx) 和 [[合同]]。\n", encoding="utf-8")
    result = extract_document(path)
    assert result.title == "项目"
    links = result.metadata.get("links", [])
    assert any(link["href"] == "budget.xlsx" for link in links)


def test_html_extracts_title_and_links(tmp_path: Path):
    path = tmp_path / "page.html"
    path.write_text(
        "<html><head><title>报价单</title></head><body>"
        "<h1>报价</h1><p>参见 <a href='budget.xlsx'>预算</a></p></body></html>",
        encoding="utf-8",
    )
    result = extract_document(path)
    assert result.title == "报价单"
    assert "报价" in result.full_text


def test_csv_extracts_rows(tmp_path: Path):
    path = tmp_path / "data.csv"
    path.write_text("项目,金额\n采购,1000\n差旅,500\n", encoding="utf-8")
    result = extract_document(path)
    assert "采购" in result.full_text
    assert result.blocks[0].location_json["row_start"] >= 1


def test_txt_extracts_lines(tmp_path: Path):
    path = tmp_path / "notes.txt"
    path.write_text("第一段\n\n第二段 合同\n", encoding="utf-8")
    result = extract_document(path)
    assert "第二段" in result.full_text

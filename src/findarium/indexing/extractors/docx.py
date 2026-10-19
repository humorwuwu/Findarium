from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .base import (
    Block,
    CorruptDocumentError,
    ExtractedDocument,
    count_words,
    truncate_text,
)

_HEADING_RE = re.compile(r"(?:heading|标题)\s*([1-9])", re.IGNORECASE)


def _iter_blocks(parent):
    if isinstance(parent, DocxDocument):
        element = parent.element.body
    else:
        element = parent._tc
    for child in element.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def _heading_level(paragraph: Paragraph) -> int | None:
    style = (paragraph.style.name if paragraph.style is not None else "") or ""
    match = _HEADING_RE.search(style)
    if match:
        return int(match.group(1))
    if style.lower() == "title":
        return 1
    try:
        level = paragraph.paragraph_format.outline_level
        if level is not None:
            return int(level) + 1
    except (AttributeError, TypeError, ValueError):
        pass
    return None


def extract(path: Path) -> ExtractedDocument:
    try:
        doc = Document(str(path))
    except Exception as exc:  # python-docx raises several low-level exceptions
        raise CorruptDocumentError(f"无法打开 DOCX：{exc}") from exc

    title = (doc.core_properties.title or "").strip() or path.stem
    author = (doc.core_properties.author or "").strip() or None
    metadata = {
        "created": str(doc.core_properties.created) if doc.core_properties.created else None,
        "modified": str(doc.core_properties.modified) if doc.core_properties.modified else None,
        "subject": doc.core_properties.subject,
        "keywords": doc.core_properties.keywords,
    }

    blocks: list[Block] = []
    heading_stack: list[str] = []
    paragraph_index = 0
    table_index = 0
    total_text: list[str] = []

    for item in _iter_blocks(doc):
        if isinstance(item, Paragraph):
            paragraph_index += 1
            text = item.text.strip()
            if not text:
                continue
            level = _heading_level(item)
            if level is not None:
                heading_stack = heading_stack[: max(0, level - 1)]
                heading_stack.append(text)
                blocks.append(
                    Block(
                        kind="heading",
                        text=text,
                        location_label=" / ".join(heading_stack),
                        location_json={
                            "paragraph": paragraph_index,
                            "heading_path": list(heading_stack),
                        },
                        heading_path=list(heading_stack),
                    )
                )
            else:
                blocks.append(
                    Block(
                        kind="paragraph",
                        text=text,
                        location_label=" / ".join(heading_stack)
                        or f"段落 {paragraph_index}",
                        location_json={
                            "paragraph": paragraph_index,
                            "heading_path": list(heading_stack),
                        },
                        heading_path=list(heading_stack),
                    )
                )
            total_text.append(text)
        else:
            table_index += 1
            rows: list[str] = []
            for row in item.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                text = "\n".join(rows)
                blocks.append(
                    Block(
                        kind="table",
                        text=text,
                        location_label=(
                            " / ".join(heading_stack) if heading_stack else ""
                        )
                        + (f" · 表格 {table_index}" if heading_stack else f"表格 {table_index}"),
                        location_json={
                            "table": table_index,
                            "heading_path": list(heading_stack),
                        },
                        heading_path=list(heading_stack),
                    )
                )
                total_text.append(text)

    # Headers/footers are useful for contracts and reports. Avoid duplicating
    # identical boilerplate across many pages.
    seen_headers: set[str] = set()
    for section in doc.sections:
        for part, label in ((section.header, "页眉"), (section.footer, "页脚")):
            try:
                text = "\n".join(p.text.strip() for p in part.paragraphs if p.text.strip())
            except Exception:
                continue
            if text and text not in seen_headers:
                seen_headers.add(text)
                blocks.append(
                    Block(
                        kind=label,
                        text=text,
                        location_label=label,
                        location_json={"part": label},
                        heading_path=[label],
                    )
                )
                total_text.append(text)

    full_text = "\n".join(total_text)
    text_limited, truncated = truncate_text(full_text)
    if truncated:
        # We keep all blocks but the pipeline applies its own character limit;
        # mark partial so the UI can explain the truncation.
        warnings = ["文档内容超过单文档字符上限，索引已按上限截断"]
    else:
        warnings = []
    word_count = count_words(full_text)
    return ExtractedDocument(
        title=title,
        author=author,
        metadata=metadata,
        blocks=blocks,
        warnings=warnings,
        partial=truncated,
        word_count=word_count,
    )

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ...i18n import index_lang, msg, translate
from .base import Block, CorruptDocumentError, ExtractedDocument, count_words


def _shape_texts(shape) -> list[str]:
    texts: list[str] = []
    try:
        if shape.has_text_frame:
            text = "\n".join(
                paragraph.text.strip()
                for paragraph in shape.text_frame.paragraphs
                if paragraph.text.strip()
            )
            if text:
                texts.append(text)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    texts.append(" | ".join(cells))
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for child in shape.shapes:
                texts.extend(_shape_texts(child))
    except Exception:
        pass
    return texts


def extract(path: Path) -> ExtractedDocument:
    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise CorruptDocumentError(
            msg("error.open_failed", format="PPTX", error=str(exc))
        ) from exc

    props = presentation.core_properties
    title = (props.title or "").strip() or path.stem
    author = (props.author or "").strip() or None
    metadata = {
        "created": str(props.created) if props.created else None,
        "modified": str(props.modified) if props.modified else None,
        "subject": props.subject,
    }

    blocks: list[Block] = []
    all_text: list[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            parts.extend(_shape_texts(shape))
        try:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            notes = ""
        if notes:
            parts.append(f"{translate("text.notes", index_lang())}:\n" + notes)
        if not parts:
            continue
        slide_title = ""
        try:
            if slide.shapes.title is not None:
                slide_title = slide.shapes.title.text.strip()
        except Exception:
            pass
        text = "\n".join(parts)
        heading_path = [slide_title] if slide_title else []
        blocks.append(
            Block(
                kind="slide",
                text=text,
                location_label=translate("location.slide", index_lang(), number=number),
                location_json={"slide": number},
                heading_path=heading_path,
            )
        )
        all_text.append(f"{translate("text.slide_number", index_lang(), number=number)}\n{text}")

    full_text = "\n\n".join(all_text)
    return ExtractedDocument(
        title=title,
        author=author,
        metadata=metadata,
        blocks=blocks,
        word_count=count_words(full_text),
        page_count=len(presentation.slides),
    )

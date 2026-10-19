from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ...config import MAX_EXCEL_ROWS_PER_CHUNK, MAX_EXCEL_ROWS_PER_SHEET, MAX_TEXT_CHARS
from ...i18n import index_lang, msg, translate
from .base import Block, CorruptDocumentError, ExtractedDocument, count_words


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:g}"
    text = str(value).strip().replace("\r", " ").replace("\n", " ")
    return text[:300]


def extract(path: Path) -> ExtractedDocument:
    try:
        workbook = load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise CorruptDocumentError(
            msg("error.open_failed", format="XLSX", error=str(exc))
        ) from exc

    props = workbook.properties
    title = (props.title or "").strip() or path.stem
    author = (props.creator or "").strip() or None
    metadata = {
        "created": str(props.created) if props.created else None,
        "modified": str(props.modified) if props.modified else None,
        "subject": props.subject,
    }

    blocks: list[Block] = []
    warnings: list[str] = []
    total_chars = 0
    partial = False
    sheet_count = len(workbook.worksheets)
    all_text: list[str] = []

    for sheet in workbook.worksheets:
        rows: list[tuple[int, str]] = []
        header_text = ""
        for row_index, values in enumerate(
            sheet.iter_rows(values_only=True), start=1
        ):
            if row_index > MAX_EXCEL_ROWS_PER_SHEET:
                warnings.append(
                    msg(
                        "warn.xlsx_rows_exceeded",
                        sheet=sheet.title,
                        limit=MAX_EXCEL_ROWS_PER_SHEET,
                    )
                )
                partial = True
                break
            cells = [_cell_text(value) for value in values]
            while cells and not cells[-1]:
                cells.pop()
            if not any(cells):
                continue
            line = " | ".join(cells)
            if not header_text:
                header_text = line
            rows.append((row_index, line))

        if not rows:
            continue

        for start in range(0, len(rows), MAX_EXCEL_ROWS_PER_CHUNK):
            group = rows[start : start + MAX_EXCEL_ROWS_PER_CHUNK]
            row_start = group[0][0]
            row_end = group[-1][0]
            body = "\n".join(f"{number}: {line}" for number, line in group)
            if header_text:
                body = f"{translate("text.column_header", index_lang())}: {header_text}\n{body}"
            location_label = translate(
                "location.sheet_rows",
                index_lang(),
                sheet=sheet.title,
                start=row_start,
                end=row_end,
            )
            block = Block(
                kind="sheet",
                text=body,
                location_label=location_label,
                location_json={
                    "sheet": sheet.title,
                    "row_start": row_start,
                    "row_end": row_end,
                },
                heading_path=[sheet.title],
            )
            blocks.append(block)
            all_text.append(body)
            total_chars += len(body)
            if total_chars > MAX_TEXT_CHARS:
                partial = True
                warnings.append(msg("warn.truncated"))
                break
        if total_chars > MAX_TEXT_CHARS:
            break

    workbook.close()
    full_text = "\n".join(all_text)
    return ExtractedDocument(
        title=title,
        author=author,
        metadata=metadata,
        blocks=blocks,
        warnings=warnings,
        partial=partial,
        word_count=count_words(full_text),
        sheet_count=sheet_count,
    )
